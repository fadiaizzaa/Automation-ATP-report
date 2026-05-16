"""
Shared YAML loader and path helpers for the week ingest / pipeline scripts.

Most paths resolve relative to the YAML directory.

``pipeline.master_workbook`` and ``pipeline.combine_workbook`` resolve relative to
the config directory first. Combine also checks the master's folder for a basename
match. Lastly the filename is sought under ``{output_base}/``.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence

import yaml
from openpyxl.utils import column_index_from_string, get_column_letter

_COL_LETTER_ONLY = re.compile(r"^[A-Za-z]{1,3}$")
_TRAILING_NUMBER = re.compile(r"(\d+)\s*$")

# --- Workbook ids used in YAML ``workbooks`` ---------------------------------

WB_NETWORK_RESOURCE = "network_resource_statistics"
WB_FIBER_COMPUTED = "fiber_computed"
WB_OMS_COMPUTED = "oms_computed"
WB_OCH_COMPUTED = "och_computed"
WB_OMS_TRAIL = "oms_trail_integrity"
WB_OCH_TRAIL = "och_trail_integrity"
WB_SERVICE_ROUTING = "service_routing_report"


@dataclass(frozen=True)
class SheetSpec:
    header_row_1based: int
    columns: dict[str, str]


@dataclass(frozen=True)
class PipelinePaths:
    config_path: Path
    config_dir: Path
    week_label: str
    output_base: Path
    computed_dir: Path
    cpq_week_dir: Path
    nms_week_dir: Path
    master_workbook: Path
    combine_workbook: Path
    current_network_resource: Path
    previous_network_resource: Path
    fiber_computed_out: Path
    oms_computed_out: Path
    och_computed_out: Path
    service_routing_raw: Path
    service_routing_computed_out: Path
    oms_trail_source: Path
    och_trail_source: Path


def _label_number(label: str) -> str:
    match = _TRAILING_NUMBER.search(label.strip())
    if not match:
        raise ValueError(f"Could not find week number in label {label!r}")
    return match.group(1)


def _config_placeholder_values(
    cfg: dict[str, Any],
    week_label_override: str | None = None,
) -> dict[str, str]:
    week_label = week_label_override if week_label_override is not None else cfg.get("week_label")
    if not isinstance(week_label, str) or not week_label.strip():
        raise ValueError("week_label must be a non-empty string")
    week_label = week_label.strip()
    week_number = _label_number(week_label)

    values = {
        "week_label": week_label,
        "week_number": week_number,
        "week_short": f"W{week_number}",
    }

    for key, value in cfg.items():
        if key in values or key in {"inputs", "output_base", "paths", "pipeline", "workbooks"}:
            continue
        if isinstance(value, (str, int, float)):
            values[key] = str(value).strip()

    values.update(
        {
            "week_label": week_label,
            "week_number": week_number,
            "week_short": f"W{week_number}",
        }
    )

    previous_week_label = cfg.get("previous_week_label")
    if isinstance(previous_week_label, str) and previous_week_label.strip():
        previous_week_label = previous_week_label.strip()
        previous_week_number = _label_number(previous_week_label)
        values.update(
            {
                "previous_week_label": previous_week_label,
                "previous_week_number": previous_week_number,
                "previous_week_short": f"W{previous_week_number}",
            }
        )

    return values


def config_placeholders(
    cfg: dict[str, Any],
    week_label_override: str | None = None,
) -> dict[str, str]:
    """Return supported placeholders for reusable path strings in ingest YAML."""
    values = _config_placeholder_values(cfg, week_label_override)
    output_base = cfg.get("output_base")
    if isinstance(output_base, str) and output_base.strip():
        values["output_base"] = output_base.strip().format(**values)
    return values


def expand_config_value(raw: str, cfg: dict[str, Any]) -> str:
    """Expand placeholders like ``{week_label}`` in a config string."""
    try:
        return raw.format(**config_placeholders(cfg))
    except KeyError as e:
        key = e.args[0]
        raise ValueError(f"Unknown placeholder {{{key}}} in config value {raw!r}") from e


def expand_config_value_for_week(raw: str, cfg: dict[str, Any], week_label: str) -> str:
    """Expand a config string as if the current week were ``week_label``."""
    try:
        return raw.format(**config_placeholders(cfg, week_label_override=week_label))
    except KeyError as e:
        key = e.args[0]
        raise ValueError(f"Unknown placeholder {{{key}}} in config value {raw!r}") from e


def configured_output_base(cfg: dict[str, Any], config_dir: Path) -> Path:
    out_base_s = cfg.get("output_base", "input")
    if not isinstance(out_base_s, str):
        raise ValueError("output_base must be a string")
    return resolve_under(config_dir, expand_config_value(out_base_s, cfg))


def resolve_under(config_dir: Path, raw: str) -> Path:
    raw = (raw or "").strip()
    if not raw:
        raise ValueError("Empty path")
    p = Path(raw)
    if p.is_absolute():
        return p.resolve()
    return (config_dir / p).resolve()


def load_yaml(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as f:
        data = yaml.safe_load(f)
    if not isinstance(data, dict):
        raise ValueError("Config root must be a mapping")
    return data


def normalize_header(cell: Any) -> str:
    if cell is None:
        return ""
    return str(cell).strip()


def resolve_column_ref(headers: list[Any], yaml_val: str, key_for_errors: str) -> tuple[int, str]:
    """
    Resolve YAML column entry to (0-based index, Excel column letter).

    If ``yaml_val`` is a bare column letter (e.g. ``H``, ``AC``), use it directly
    (helps when headers duplicate the same text).
    Otherwise match ``yaml_val`` as exact header text after strip.
    """
    v = (yaml_val or "").strip()
    if not v:
        raise ValueError(f"Column key {key_for_errors!r}: empty value")
    if _COL_LETTER_ONLY.fullmatch(v):
        letter = v.upper()
        idx = column_index_from_string(letter) - 1
        return idx, letter
    idx = column_index_from_headers(headers, v, key_for_errors)
    return idx, get_column_letter(idx + 1)


def column_index_from_headers(headers: list[Any], title: str, key_for_errors: str) -> int:
    """
    Return 0-based column index for ``title`` (exact match after strip).
    Raises ValueError if missing or ambiguous.
    """
    want = title.strip()
    matches: list[int] = []
    for i, h in enumerate(headers):
        if normalize_header(h) == want:
            matches.append(i)
    if not matches:
        preview = ", ".join(
            f"{i}:{normalize_header(h)!r}" for i, h in enumerate(headers[:40]) if normalize_header(h)
        )
        raise ValueError(f"Column key {key_for_errors!r}: header {want!r} not found. Seen: [{preview}]")
    if len(matches) > 1:
        raise ValueError(f"Column key {key_for_errors!r}: duplicate header {want!r} at indices {matches}")
    return matches[0]


def get_sheet_spec(workbooks: dict[str, Any], wb_key: str, sheet_name: str) -> SheetSpec:
    wb = workbooks.get(wb_key)
    if not isinstance(wb, dict):
        raise KeyError(f"workbooks.{wb_key} is missing or not a mapping")
    sheets = wb.get("sheets")
    if not isinstance(sheets, dict):
        raise KeyError(f"workbooks.{wb_key}.sheets is missing or not a mapping")
    block = sheets.get(sheet_name)
    if not isinstance(block, dict):
        raise KeyError(f"workbooks.{wb_key}.sheets[{sheet_name!r}] missing")
    hr = block.get("header_row")
    if not isinstance(hr, int) or hr < 1:
        raise ValueError(f"{wb_key}/{sheet_name}: header_row must be integer >= 1")
    cols = block.get("columns")
    if cols is None:
        cols = {}
    if not isinstance(cols, dict):
        raise ValueError(f"{wb_key}/{sheet_name}: columns must be a mapping")
    col_map = {str(k): str(v) for k, v in cols.items()}
    return SheetSpec(header_row_1based=hr, columns=col_map)


def idx_map_from_headers(headers: list[Any], columns_yaml: dict[str, str], keys: list[str]) -> dict[str, int]:
    """Resolve listed internal keys to 0-based indices (supports bare column letters in YAML)."""
    out: dict[str, int] = {}
    for k in keys:
        if k not in columns_yaml:
            raise KeyError(f"Missing columns.{k} in YAML")
        out[k] = resolve_column_ref(headers, columns_yaml[k], k)[0]
    return out


def try_resolve_nms_workbook(folder: Path, week_label: str, stem_suffix: str) -> Path | None:
    """Prefer .xlsm then .xlsx for ``{week_label}_{stem_suffix}``."""
    for ext in (".xlsm", ".xlsx"):
        p = folder / f"{week_label}_{stem_suffix}{ext}"
        if p.is_file():
            return p.resolve()
    return None


def try_resolve_workbook_path(
    cfg_dir: Path,
    output_base: Path,
    week_label: str,
    configured: str,
    extra_search_roots: Sequence[Path] | None = None,
) -> Path:
    """
    Resolve ``pipeline.master_workbook`` / ``combine_workbook``.

    1. Path relative to the ingest YAML directory (or absolute).
    2. If missing and ``extra_search_roots`` is given, try ``{root}/{basename}`` for each root
       (used so a bare combine filename resolves next to master when master lives elsewhere).
    3. Else ``{output_base}/{basename}``.
    """
    primary = resolve_under(cfg_dir, configured)
    name = Path(configured).name
    if primary.is_file():
        return primary.resolve()
    for root in extra_search_roots or []:
        alt = (root.resolve() / name).resolve()
        if alt.is_file():
            return alt
    under_output_base = (output_base / name).resolve()
    if under_output_base.is_file():
        return under_output_base
    return primary.resolve()


def resolve_pipeline_output_path(
    cfg_dir: Path,
    computed_dir: Path,
    cfg: dict[str, Any],
    raw: str | None,
    default_name: str,
) -> Path:
    if isinstance(raw, str) and raw.strip():
        expanded = expand_config_value(raw.strip(), cfg)
        configured_path = Path(expanded)
        if configured_path.name == expanded and not configured_path.is_absolute():
            return computed_dir / expanded
        return resolve_under(cfg_dir, expanded)
    return computed_dir / default_name


def previous_output_base(output_base: Path, week_label: str, previous_week_label: str) -> Path:
    parts = output_base.parts
    for idx, part in enumerate(parts):
        if part.lower() == week_label.lower():
            return Path(*parts[:idx], previous_week_label, *parts[idx + 1 :])
    return output_base / previous_week_label


def prepared_input_root(output_base: Path) -> Path:
    if output_base.name.lower() == "input":
        return output_base
    return output_base / "input"


def build_pipeline_paths(cfg: dict[str, Any], config_path: Path) -> PipelinePaths:
    cfg_dir = config_path.parent.resolve()
    week_label = cfg.get("week_label")
    if not isinstance(week_label, str) or not week_label.strip():
        raise ValueError("week_label must be a non-empty string")

    output_base = configured_output_base(cfg, cfg_dir)
    input_root = prepared_input_root(output_base)

    pipe = cfg.get("pipeline")
    if not isinstance(pipe, dict):
        raise ValueError("pipeline mapping is required")

    computed_s = pipe.get("computed_dir", "computed")
    if not isinstance(computed_s, str):
        raise ValueError("pipeline.computed_dir must be a string")
    computed_s = expand_config_value(computed_s, cfg)
    computed_dir = resolve_under(cfg_dir, computed_s)

    master_s = pipe.get("master_workbook")
    combine_s = pipe.get("combine_workbook")
    if not isinstance(master_s, str) or not isinstance(combine_s, str):
        raise ValueError("pipeline.master_workbook and combine_workbook must be strings")
    master_s = expand_config_value(master_s, cfg)
    combine_s = expand_config_value(combine_s, cfg)

    cpq_week = input_root / "CPQ"
    nms_week = input_root / "NMS"

    cur_nrr = try_resolve_nms_workbook(cpq_week, week_label, "Network Resource Statistics")
    if cur_nrr is None:
        raise FileNotFoundError(f"Network Resource Statistics not found under {cpq_week}")

    prev_path: Path | None = None
    prev_wb_s = pipe.get("previous_network_resource_workbook")
    if isinstance(prev_wb_s, str) and prev_wb_s.strip():
        prev_path = resolve_under(cfg_dir, expand_config_value(prev_wb_s.strip(), cfg))
    else:
        prev_label = cfg.get("previous_week_label")
        if isinstance(prev_label, str) and prev_label.strip():
            prev_cpq = prepared_input_root(
                previous_output_base(output_base, week_label, prev_label.strip())
            ) / "CPQ"
            prev_path = try_resolve_nms_workbook(prev_cpq, prev_label.strip(), "Network Resource Statistics")

    if prev_path is None:
        raise ValueError("Set previous_week_label or pipeline.previous_network_resource_workbook")

    svc_out_name = pipe.get("service_routing_computed")
    svc_out = resolve_pipeline_output_path(
        cfg_dir,
        computed_dir,
        cfg,
        svc_out_name,
        f"{week_label}_ServiceRouting_computed2.xlsx",
    )

    oms_src = cpq_week / f"{week_label}_OMS Trail Integrity.xlsx"
    och_src = cpq_week / f"{week_label}_OCh Trail Integrity.xlsx"
    if not oms_src.is_file():
        raise FileNotFoundError(str(oms_src))
    if not och_src.is_file():
        raise FileNotFoundError(str(och_src))

    svc_raw = cpq_week / f"{week_label}_ServiceRoutingReport.xlsx"
    if not svc_raw.is_file():
        raise FileNotFoundError(str(svc_raw))

    master_workbook = try_resolve_workbook_path(cfg_dir, output_base, week_label, master_s)
    combine_workbook = try_resolve_workbook_path(
        cfg_dir,
        output_base,
        week_label,
        combine_s,
        extra_search_roots=[master_workbook.parent],
    )

    return PipelinePaths(
        config_path=config_path.resolve(),
        config_dir=cfg_dir,
        week_label=week_label.strip(),
        output_base=output_base,
        computed_dir=computed_dir,
        cpq_week_dir=cpq_week,
        nms_week_dir=nms_week,
        master_workbook=master_workbook,
        combine_workbook=combine_workbook,
        current_network_resource=cur_nrr,
        previous_network_resource=prev_path,
        fiber_computed_out=computed_dir / f"{week_label.strip()}_Fiber_computed.xlsx",
        oms_computed_out=computed_dir / f"{week_label.strip()}_OMS_computed.xlsx",
        och_computed_out=computed_dir / f"{week_label.strip()}_OCh_computed.xlsx",
        service_routing_raw=svc_raw,
        service_routing_computed_out=svc_out,
        oms_trail_source=oms_src,
        och_trail_source=och_src,
    )
